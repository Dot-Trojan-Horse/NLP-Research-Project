#!/usr/bin/env python3
"""Run adapted Allosaurus independently on all Assignment 3 study streets.

Outputs:
    study_phone_recognition.csv
        Raw structured recognition results.

    study_phone_recognition_human_readable.xlsx
        Human-readable workbook with:
        - Recognition Results
        - Street Summary
        - Problem Records
        - Method Notes

    study_phone_recognition_report.txt
        Human-readable terminal-style research report.

    study_phone_recognition_summary.csv
        One row per street.

    study_phone_recognition_problems.csv
        Any missing, empty, or failed recognition records.

IMPORTANT:
    Allosaurus outputs are independent phone-recognition hypotheses.
    They are NOT automatically treated as ground truth.
    Ground-truth validation is handled separately.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from allosaurus.app import read_recognizer
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]

DEFAULT_REGISTRY = (
    PROJECT_ROOT / "results" / "assignment3" / "ground_truth"
    / "assignment3_ground_truth_master.csv"
)

DEFAULT_GT = PROJECT_ROOT / "recordings" / "ground_truth"
DEFAULT_SYSTEM = PROJECT_ROOT / "recordings" / "wav_system"

DEFAULT_OUTPUT = (
    PROJECT_ROOT / "results" / "assignment3" / "phone_recognition"
)


def variants(street: str) -> list[str]:
    return [
        street,
        street.replace("_", " "),
        street.replace(" ", "_"),
    ]


def resolve_gt(street: str) -> Path | None:
    for variant in variants(street):
        path = DEFAULT_GT / variant
        if path.is_dir():
            return path
    return None


def resolve_system(street: str) -> Path | None:
    for variant in variants(street):
        path = DEFAULT_SYSTEM / f"{variant}.wav"
        if path.is_file():
            return path
    return None


def tokenize(text: str) -> list[str]:
    return [token for token in text.strip().split() if token]


def recognize_one(recognizer, wav: Path, lang: str) -> tuple[str, str]:
    if not wav.is_file():
        return "missing_audio", ""

    try:
        recognized_text = recognizer.recognize(
            str(wav),
            lang_id=lang,
        )

        recognized = tokenize(recognized_text)

        if recognized:
            return "ok", " ".join(recognized)

        return "empty_recognition", ""

    except Exception as exc:
        return f"recognition_error: {exc}", ""


def write_xlsx(final_rows, summary_rows, problem_rows, output_file):
    """Create a human-readable Excel workbook."""

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for the Excel report. "
            "Install it with: pip install openpyxl"
        ) from exc

    wb = Workbook()

    # ---------------------------------------------------------------
    # Recognition Results
    # ---------------------------------------------------------------
    ws = wb.active
    ws.title = "Recognition Results"

    headers = [
        "Street Name",
        "Recording Type",
        "Recording File",
        "Ground-Truth Status",
        "Reference Phones",
        "Recognition Status",
        "Recognised Phones",
        "Phone Count",
    ]

    ws.append(headers)

    for row in final_rows:
        ws.append([
            row["street"].replace("_", " "),
            (
                "Native Speaker"
                if row["recording_type"] == "native"
                else "Apple Maps"
            ),
            Path(row["recording"]).name,
            row["validation_status"],
            row["reference_phones"],
            row["status"],
            row["recognized_phones"],
            int(row["phone_count"]),
        ])

    style_sheet(ws, headers, [28, 20, 20, 25, 24, 24, 55, 12])

    if ws.max_row > 1:
        ref = f"A1:H{ws.max_row}"
        tab = Table(displayName="RecognitionResults", ref=ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tab)

    # ---------------------------------------------------------------
    # Street Summary
    # ---------------------------------------------------------------
    sws = wb.create_sheet("Street Summary")

    summary_headers = list(summary_rows[0].keys()) if summary_rows else [
        "Street"
    ]

    sws.append(summary_headers)

    for row in summary_rows:
        sws.append(list(row.values()))

    style_sheet(
        sws,
        summary_headers,
        [28, 18, 18, 21, 22, 16, 23, 16, 18, 18, 25],
    )

    if sws.max_row > 1:
        ref = f"A1:{get_column_letter(sws.max_column)}{sws.max_row}"
        tab = Table(displayName="StreetSummary", ref=ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
        )
        sws.add_table(tab)

    # ---------------------------------------------------------------
    # Problem Records
    # ---------------------------------------------------------------
    pws = wb.create_sheet("Problem Records")

    problem_headers = headers
    pws.append(problem_headers)

    for row in problem_rows:
        pws.append([
            row["street"].replace("_", " "),
            (
                "Native Speaker"
                if row["recording_type"] == "native"
                else "Apple Maps"
            ),
            Path(row["recording"]).name,
            row["validation_status"],
            row["reference_phones"],
            row["status"],
            row["recognized_phones"],
            int(row["phone_count"]),
        ])

    style_sheet(
        pws,
        problem_headers,
        [28, 20, 20, 25, 24, 30, 55, 12],
    )

    if pws.max_row > 1:
        ref = f"A1:H{pws.max_row}"
        tab = Table(displayName="ProblemRecords", ref=ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
        )
        pws.add_table(tab)

    # ---------------------------------------------------------------
    # Method Notes
    # ---------------------------------------------------------------
    nws = wb.create_sheet("Method Notes")

    notes = [
        ("Item", "Explanation"),
        (
            "Purpose",
            "Independent phone recognition of the native-speaker and "
            "Apple Maps study recordings."
        ),
        (
            "Recognizer",
            "Adapted Allosaurus model trained/adapted using the isiZulu "
            "NCHLT preparation."
        ),
        (
            "Ground truth",
            "Recognised phone sequences are independent model hypotheses "
            "and are not automatically treated as ground truth."
        ),
        (
            "Validation",
            "Validated reference sequences must be established separately "
            "using the project's ground-truth procedure."
        ),
        (
            "Speaker variation",
            "Differences between native speakers should not automatically "
            "be classified as pronunciation errors."
        ),
        (
            "Street suffixes",
            "Original recordings are retained. No original audio is "
            "trimmed, overwritten, or deleted by this script."
        ),
        (
            "PER",
            "PER and Levenshtein analysis should be performed only after "
            "an appropriate reference sequence has been validated."
        ),
    ]

    for row in notes:
        nws.append(row)

    nws.column_dimensions["A"].width = 25
    nws.column_dimensions["B"].width = 105

    for cell in nws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in nws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    nws.freeze_panes = "A2"
    nws.row_dimensions[1].height = 30

    wb.save(output_file)


def style_sheet(ws, headers, widths):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    bottom = Side(style="thin", color="D9E2F3")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(bottom=bottom)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34


def write_csv(path: Path, rows: list[dict], fields: list[str]):
    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=fields,
        )
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    ap = argparse.ArgumentParser(
        description=(
            "Recognise phone sequences independently for all study "
            "street recordings and create research-friendly outputs."
        )
    )

    ap.add_argument(
        "--model",
        required=True,
        help="Allosaurus model name, e.g. zul_nchlt_model",
    )

    ap.add_argument(
        "--lang",
        default="zul",
        help="Allosaurus language/inventory ID",
    )

    ap.add_argument(
        "--registry",
        type=Path,
        default=DEFAULT_REGISTRY,
    )

    ap.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
    )

    ap.add_argument(
        "--streets",
        nargs="*",
        help="Optional street keys. Omit to process all streets.",
    )

    args = ap.parse_args()

    print("=" * 78)
    print("ISI ZULU STREET-NAME PHONE RECOGNITION")
    print("=" * 78)
    print()
    print(f"Project root      : {PROJECT_ROOT}")
    print(f"Ground-truth root : {DEFAULT_GT}")
    print(f"Apple Maps root   : {DEFAULT_SYSTEM}")
    print(f"Registry          : {args.registry}")
    print(f"Output directory  : {args.output.resolve()}")
    print()

    if not args.registry.is_file():
        raise FileNotFoundError(
            f"Reference registry not found:\n{args.registry}"
        )

    args.output.mkdir(
        parents=True,
        exist_ok=True,
    )

    with args.registry.open(
        encoding="utf-8-sig",
        newline="",
    ) as f:
        registry = list(csv.DictReader(f))

    wanted = set(args.streets) if args.streets else None

    print(f"Loading Allosaurus model: {args.model}")

    recognizer = read_recognizer(args.model)

    if (
        args.lang != "ipa"
        and not recognizer.is_available(args.lang)
    ):
        raise ValueError(
            f"Language inventory '{args.lang}' is not available "
            f"in model '{args.model}'."
        )

    print("Allosaurus model loaded.")
    print()

    # ---------------------------------------------------------------
    # Discover recordings
    # ---------------------------------------------------------------

    rows = []

    for r in registry:
        street = r.get("street", "").strip()

        if not street:
            continue

        if wanted and street not in wanted:
            continue

        validation_status = r.get(
            "validation_status",
            "",
        )

        reference_phones = r.get(
            "ground_truth_phone_sequence",
            "",
        )

        print(f"Street: {street.replace('_', ' ')}")

        gt_dir = resolve_gt(street)

        if gt_dir:
            native_wavs = sorted(gt_dir.glob("*.wav"))

            print(
                f"  Native directory : FOUND"
            )
            print(
                f"  Native WAV files : {len(native_wavs)}"
            )

            for wav in native_wavs:
                print(
                    f"    + native: {wav.name}"
                )

                rows.append({
                    "street": street,
                    "recording_type": "native",
                    "recording": str(wav.resolve()),
                    "validation_status": validation_status,
                    "reference_phones": reference_phones,
                    "wav_exists": "yes",
                    "wav_size_bytes": wav.stat().st_size,
                })
        else:
            print("  Native directory : NOT FOUND")

        system = resolve_system(street)

        if system:
            print("  Apple Maps WAV   : FOUND")
            print(f"    + apple: {system.name}")

            rows.append({
                "street": street,
                "recording_type": "apple_maps",
                "recording": str(system.resolve()),
                "validation_status": validation_status,
                "reference_phones": reference_phones,
                "wav_exists": "yes",
                "wav_size_bytes": system.stat().st_size,
            })
        else:
            print("  Apple Maps WAV   : NOT FOUND")

        print()

    print("=" * 78)
    print("RECORDING DISCOVERY")
    print("=" * 78)

    native_count = sum(
        r["recording_type"] == "native"
        for r in rows
    )

    apple_count = sum(
        r["recording_type"] == "apple_maps"
        for r in rows
    )

    print(f"Total recordings discovered : {len(rows)}")
    print(f"Native recordings           : {native_count}")
    print(f"Apple Maps recordings       : {apple_count}")
    print()

    if not rows:
        raise RuntimeError(
            "No recordings were discovered."
        )

    # ---------------------------------------------------------------
    # Recognition
    # ---------------------------------------------------------------

    print("=" * 78)
    print("RUNNING ALLOSAURUS RECOGNITION")
    print("=" * 78)
    print()

    final_rows = []

    for number, row in enumerate(rows, start=1):
        street = row["street"]
        rec_type = row["recording_type"]
        wav = Path(row["recording"])

        display_type = (
            "Native Speaker"
            if rec_type == "native"
            else "Apple Maps"
        )

        print(
            f"[{number}/{len(rows)}] "
            f"{street.replace('_', ' ')} | "
            f"{display_type} | {wav.name}"
        )

        status, phone_string = recognize_one(
            recognizer,
            wav,
            args.lang,
        )

        phones = tokenize(phone_string)

        final_rows.append({
            "street": street,
            "recording_type": rec_type,
            "recording": str(wav),
            "validation_status": row["validation_status"],
            "reference_phones": row["reference_phones"],
            "status": status,
            "recognized_phones": phone_string,
            "phone_count": len(phones),
        })

        if status == "ok":
            print(f"    Phones : {phone_string}")
            print(f"    Count  : {len(phones)}")
        else:
            print(f"    Status : {status}")

        print()

    # ---------------------------------------------------------------
    # Statistics
    # ---------------------------------------------------------------

    total = len(final_rows)

    successful = sum(
        r["status"] == "ok"
        for r in final_rows
    )

    missing = sum(
        r["status"] == "missing_audio"
        for r in final_rows
    )

    empty = sum(
        r["status"] == "empty_recognition"
        for r in final_rows
    )

    errors = sum(
        str(r["status"]).startswith("recognition_error:")
        for r in final_rows
    )

    # ---------------------------------------------------------------
    # Street-level summary
    # ---------------------------------------------------------------

    streets = sorted(
        {r["street"] for r in final_rows}
    )

    summary_rows = []

    for street in streets:
        sr = [
            r for r in final_rows
            if r["street"] == street
        ]

        nr = [
            r for r in sr
            if r["recording_type"] == "native"
        ]

        ar = [
            r for r in sr
            if r["recording_type"] == "apple_maps"
        ]

        summary_rows.append({
            "Street": street.replace("_", " "),
            "Native recordings": len(nr),
            "Native successful": sum(
                r["status"] == "ok"
                for r in nr
            ),
            "Apple Maps recordings": len(ar),
            "Apple Maps successful": sum(
                r["status"] == "ok"
                for r in ar
            ),
            "Total recordings": len(sr),
            "Successful recognitions": sum(
                r["status"] == "ok"
                for r in sr
            ),
            "Missing audio": sum(
                r["status"] == "missing_audio"
                for r in sr
            ),
            "Empty recognitions": sum(
                r["status"] == "empty_recognition"
                for r in sr
            ),
            "Recognition errors": sum(
                str(r["status"]).startswith(
                    "recognition_error:"
                )
                for r in sr
            ),
            "Ground-truth status": (
                sr[0]["validation_status"]
                if sr
                else ""
            ),
        })

    problem_rows = [
        r for r in final_rows
        if r["status"] != "ok"
    ]

    # ---------------------------------------------------------------
    # Write outputs
    # ---------------------------------------------------------------

    fields = [
        "street",
        "recording_type",
        "recording",
        "validation_status",
        "reference_phones",
        "status",
        "recognized_phones",
        "phone_count",
    ]

    csv_file = (
        args.output
        / "study_phone_recognition.csv"
    )

    summary_file = (
        args.output
        / "study_phone_recognition_summary.csv"
    )

    problems_file = (
        args.output
        / "study_phone_recognition_problems.csv"
    )

    report_file = (
        args.output
        / "study_phone_recognition_report.txt"
    )

    xlsx_file = (
        args.output
        / "study_phone_recognition_human_readable.xlsx"
    )

    write_csv(
        csv_file,
        final_rows,
        fields,
    )

    write_csv(
        summary_file,
        summary_rows,
        list(summary_rows[0].keys()),
    )

    write_csv(
        problems_file,
        problem_rows,
        fields,
    )

    # ---------------------------------------------------------------
    # Human-readable TXT report
    # ---------------------------------------------------------------

    report = []

    report.extend([
        "=" * 78,
        "ISI ZULU STREET-NAME PHONE RECOGNITION REPORT",
        "=" * 78,
        "",
        "This report records independent phone hypotheses generated by",
        "the adapted Allosaurus model for the Assignment 3 study data.",
        "",
        f"Model                  : {args.model}",
        f"Language inventory     : {args.lang}",
        f"Study streets          : {len(streets)}",
        f"Total recordings      : {total}",
        f"Native recordings      : {native_count}",
        f"Apple Maps recordings  : {apple_count}",
        "",
        "IMPORTANT METHODOLOGICAL NOTE",
        "-" * 78,
        "The recognised phone sequences are independent model hypotheses.",
        "They are NOT automatically treated as ground truth.",
        "Native-speaker variation should not automatically be classified",
        "as recognition error.",
        "Ground-truth validation is performed separately.",
        "",
        "=" * 78,
        "OVERALL RESULTS",
        "=" * 78,
        "",
        f"Successful recognitions : {successful}/{total}",
        f"Missing audio           : {missing}",
        f"Empty recognitions      : {empty}",
        f"Recognition errors      : {errors}",
        "",
    ])

    for index, street in enumerate(streets, start=1):
        sr = [
            r for r in final_rows
            if r["street"] == street
        ]

        report.extend([
            "=" * 78,
            f"{index:02d}. {street.replace('_', ' ')}",
            "=" * 78,
            f"Ground-truth status: {sr[0]['validation_status']}",
            "",
            "NATIVE-SPEAKER RECOGNITIONS",
            "-" * 78,
        ])

        nr = [
            r for r in sr
            if r["recording_type"] == "native"
        ]

        if nr:
            for r in nr:
                report.append(
                    f"{Path(r['recording']).name:<10} | "
                    f"{r['status']:<20} | "
                    f"{r['recognized_phones']}"
                )
        else:
            report.append("No native recordings discovered.")

        report.extend([
            "",
            "APPLE MAPS RECOGNITION",
            "-" * 78,
        ])

        ar = [
            r for r in sr
            if r["recording_type"] == "apple_maps"
        ]

        if ar:
            for r in ar:
                report.append(
                    f"{Path(r['recording']).name:<18} | "
                    f"{r['status']:<20} | "
                    f"{r['recognized_phones']}"
                )
        else:
            report.append("No Apple Maps recording discovered.")

        report.extend([
            "",
            "REVIEW NOTE",
            "-" * 78,
            "Compare native outputs for speaker variation and recogniser",
            "consistency. Compare the Apple Maps output separately.",
            "Do not assign PER until an appropriate reference sequence",
            "has been validated.",
            "",
        ])

    report.extend([
        "=" * 78,
        "FILES CREATED",
        "=" * 78,
        "",
        str(csv_file.resolve()),
        str(summary_file.resolve()),
        str(problems_file.resolve()),
        str(report_file.resolve()),
        str(xlsx_file.resolve()),
        "",
    ])

    report_file.write_text(
        "\n".join(report),
        encoding="utf-8",
    )

    # Excel is a presentation layer over the same results.
    write_xlsx(
        final_rows,
        summary_rows,
        problem_rows,
        xlsx_file,
    )

    # ---------------------------------------------------------------
    # Final terminal summary
    # ---------------------------------------------------------------

    print("=" * 78)
    print("PHONE RECOGNITION COMPLETE")
    print("=" * 78)
    print()
    print(f"Study streets             : {len(streets)}")
    print(f"Recordings/rows written   : {total}")
    print(f"Successful recognitions   : {successful}")
    print(f"Missing audio             : {missing}")
    print(f"Empty recognitions        : {empty}")
    print(f"Recognition errors        : {errors}")
    print()
    print("Research-friendly outputs:")
    print(f"CSV     : {csv_file.resolve()}")
    print(f"Summary  : {summary_file.resolve()}")
    print(f"Problems : {problems_file.resolve()}")
    print(f"Report   : {report_file.resolve()}")
    print(f"Excel    : {xlsx_file.resolve()}")
    print()


if __name__ == "__main__":
    main()
